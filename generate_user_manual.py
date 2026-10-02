import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def create_user_manual():
    doc = Document()
    
    # Custom Styles
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    # --- COVER PAGE ---
    doc.add_heading('USER MANUAL\nWEB PENGELOLAAN BANK SAMPAH DIGITAL\n(ECOTRACK)', 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('\nDisusun untuk melengkapi dokumentasi project aplikasi manajemen sampah dan e-wallet.\n', style='Normal').alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Nama Anda / Kelompok')
    run.font.size = Pt(14)
    run.font.bold = True
    
    doc.add_page_break()

    # --- FILOSOFI WARNA ---
    doc.add_heading('Filosofi Warna', level=1)
    colors = [
        ("Emerald Green (#10B981)", "Melambangkan alam, kelestarian lingkungan, dan regenerasi kehidupan. Warna ini merepresentasikan visi utama untuk mewujudkan lingkungan yang bersih dan berkelanjutan."),
        ("Slate Dark (#0F172A)", "Melambangkan ketegasan, profesionalisme, dan kestabilan sistem. Memberikan kenyamanan visual saat membaca data, sekaligus mencerminkan keamanan data transaksi."),
        ("Frost White (#F8FAFC)", "Melambangkan kebersihan, transparansi tata kelola, dan kesederhanaan. Menghadirkan ruang antarmuka yang bersih (clean space)."),
        ("Amber Gold (#F59E0B)", "Melambangkan nilai ekonomi, semangat, dan penghargaan atas kontribusi. Menggambarkan bahwa sampah yang dipilah memiliki nilai tukar nyata berupa poin e-Wallet."),
        ("Rose Red (#F43F5E)", "Melambangkan kewaspadaan dan kehati-hatian. Digunakan pada notifikasi validasi dan aksi penghapusan (delete) data.")
    ]
    for color, desc in colors:
        doc.add_paragraph(f"{color}", style='Heading 3')
        doc.add_paragraph(desc)

    doc.add_page_break()

    # --- KEBUTUHAN SUMBER DAYA ---
    doc.add_heading('Kebutuhan Sumber Daya', level=1)
    doc.add_paragraph("Aplikasi EcoTrack berbasis web dan dapat diakses dari berbagai perangkat dengan spesifikasi minimum berikut:")
    
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Komponen'
    hdr_cells[1].text = 'Spesifikasi Minimum'
    
    specs = [
        ('Sistem Operasi', 'Windows 10/11, macOS, Linux, atau Android/iOS'),
        ('Web Browser', 'Google Chrome, Microsoft Edge, Safari, atau Mozilla Firefox'),
        ('Processor', 'Minimum Dual Core 2.0 GHz atau setara'),
        ('Memori (RAM)', 'Minimum 2 GB (Disarankan 4 GB)'),
        ('Koneksi Internet', 'Jaringan Wi-Fi / Data Seluler (Minimal 512 Kbps)'),
        ('Hak Akses User', 'Akun Warga (User) atau Akun Petugas (Admin)')
    ]
    for comp, spec in specs:
        row_cells = table.add_row().cells
        row_cells[0].text = comp
        row_cells[1].text = spec

    doc.add_page_break()

    # --- MEMULAI APLIKASI ---
    doc.add_heading('1. Memulai Aplikasi', level=1)
    
    # Login
    doc.add_heading('Halaman Login', level=2)
    if os.path.exists('audit_desktop_login.png'):
        doc.add_picture('audit_desktop_login.png', width=Inches(6.0))
    doc.add_paragraph("Pada halaman login, pengguna dapat memasukkan Email dan Password. Terdapat pemisahan akses otomatis: akun dengan hak akses admin akan diarahkan ke Admin Panel, sedangkan warga akan diarahkan ke Dasbor Warga.")
    
    # Register
    doc.add_heading('Halaman Register', level=2)
    if os.path.exists('audit_desktop_register.png'):
        doc.add_picture('audit_desktop_register.png', width=Inches(6.0))
    doc.add_paragraph("Warga yang belum memiliki akun dapat melakukan registrasi dengan mengisi Nama Lengkap, Email, dan Password. Sistem akan mengenkripsi kata sandi secara otomatis.")

    doc.add_page_break()

    # --- WARGA PANEL ---
    doc.add_heading('2. Warga Panel (User)', level=1)
    
    # Dashboard Warga
    doc.add_heading('Dashboard Warga', level=2)
    if os.path.exists('audit_desktop_user_dashboard_clean.png'):
        doc.add_picture('audit_desktop_user_dashboard_clean.png', width=Inches(6.0))
    doc.add_paragraph("Pada halaman Dashboard Warga, pengguna disambut dengan ringkasan informasi akun. Terdapat indikator Saldo Poin yang dapat ditukarkan, serta tabel ringkasan 5 transaksi terakhir. Pengguna juga akan melihat Popup interaktif yang memandu cara menukar sampah menjadi uang.")
    
    # Lapor Sampah
    doc.add_heading('Form Lapor Sampah', level=2)
    if os.path.exists('audit_desktop_form_lapor.png'):
        doc.add_picture('audit_desktop_form_lapor.png', width=Inches(6.0))
    doc.add_paragraph("Komponen / Elemen pada Halaman:\n"
                      "1. Input Berat Sampah: Warga memasukkan berat sampah dalam kilogram.\n"
                      "2. Upload Foto: Bukti fisik sampah wajib diunggah.\n"
                      "3. Catatan: Informasi tambahan lokasi/jenis sampah detail.\n"
                      "4. Tombol Kirim Laporan: Mengirim data ke antrean verifikasi admin.")

    # e-Wallet
    doc.add_heading('E-Wallet & Penukaran Poin', level=2)
    if os.path.exists('audit_desktop_ewallet.png'):
        doc.add_picture('audit_desktop_ewallet.png', width=Inches(6.0))
    doc.add_paragraph("Warga dapat mencairkan poin menjadi saldo E-Wallet:\n"
                      "1. Opsi E-Wallet: Pilih antara GoPay, DANA, OVO, atau ShopeePay.\n"
                      "2. Input Nomor HP: Nomor tujuan transfer saldo.\n"
                      "3. Jumlah Poin: Sistem akan menampilkan estimasi Rupiah secara real-time (1 Poin = Rp 100).\n"
                      "4. Riwayat: Melihat status pencairan (Menunggu, Berhasil, atau Gagal/Refund).")

    doc.add_page_break()

    # --- ADMIN PANEL ---
    doc.add_heading('3. Admin Panel', level=1)
    
    # Dashboard Admin & Manajemen Laporan
    doc.add_heading('Dashboard Admin & Verifikasi Laporan', level=2)
    if os.path.exists('audit_desktop_admin_top.png'):
        doc.add_picture('audit_desktop_admin_top.png', width=Inches(6.0))
    doc.add_paragraph("Admin Panel memberikan visibilitas penuh atas seluruh aktivitas warga:\n"
                      "1. Grafik Statistik: Pantau tren setoran mingguan/bulanan.\n"
                      "2. Tabel Laporan: Admin meninjau laporan sampah baru.\n"
                      "3. Aksi Verifikasi: Admin dapat mengubah status laporan menjadi 'Selesai'. Saat status selesai, poin akan otomatis ditambahkan ke saldo warga.")

    # Verifikasi Penarikan E-Wallet
    doc.add_heading('Verifikasi Pencairan e-Wallet', level=2)
    if os.path.exists('audit_desktop_admin_bottom.png'):
        doc.add_picture('audit_desktop_admin_bottom.png', width=Inches(6.0))
    doc.add_paragraph("Bagian Manajemen Keuangan (di bawah Dashboard Admin):\n"
                      "1. Daftar antrean pencairan e-Wallet warga (GoPay/DANA/OVO).\n"
                      "2. Klik 'Setujui' jika transfer telah dilakukan (Nomor referensi unik ECO-XXX akan diterbitkan).\n"
                      "3. Klik 'Tolak' jika nomor e-Wallet salah. Sistem akan otomatis melakukan refund poin kembali ke saldo warga.")

    doc.save('User_Manual_EcoTrack.docx')
    print("User_Manual_EcoTrack.docx successfully generated.")

if __name__ == '__main__':
    create_user_manual()
